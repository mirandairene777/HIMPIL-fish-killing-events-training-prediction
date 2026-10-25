"""
HIMPIL Phase 3a - clean the environmental workbooks.

This script preserves the source resolution. It does NOT:

* interpolate missing values;
* treat a missing symbol as zero;
* turn an annual average into a daily observation;
* merge annual water quality with monthly weather as if the timestamps matched.

Outputs (data/processed)
------------------------
water_quality_long.csv
water_quality_all_annual_wide.csv
water_quality_lake_annual.csv
station_registry.csv
water_quality_coverage_summary.csv
lake_station_year_coverage.csv
rainfall_monthly.csv
wind_monthly.csv
weather_monthly_wide.csv
weather_station_registry.csv
weather_coverage_summary.csv

Report: reports/tables/cleaning_report.md

Run from the project root:
    .venv\\Scripts\\python.exe python\\clean_data.py
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from common import (
    MONTHS,
    MONTH_NUMBER,
    PROCESSED_DIR,
    RAW_DIR,
    TABLE_DIR,
    VARIABLES,
    WEATHER_SPECS,
    as_float,
    ensure_dirs,
    find_workbook,
    is_number,
    missing_kind,
    normalise_name,
    text,
    write_csv,
    year_from_header,
)

WQ_SOURCE_IDS = {
    "nitrate": "S02",
    "phosphate": "S03",
    "bod": "S04",
    "ph": "S05",
    "temperature": "S06",
    "do": "S07",
}
WEATHER_SOURCE_IDS = {"rainfall": "S01", "wind": "S08"}

QUALITY_LABELS = {
    "blank": "blank",
    "station_discontinued": "station_discontinued",
    "no_data_symbol": "no_data",
    "ellipsis_undefined": "undefined_ellipsis",
    "other_non_numeric": "unexpected_text",
}

SECTION_LABELS = {
    "laguna de bay",
    "major tributaries",
    "ncr",
    "laguna",
    "rizal",
    "region iv-a",
    "region iv-a (calabarzon)",
    "calabarzon",
}


def find_water_header(ws) -> tuple[int, dict[int, int]] | tuple[None, dict]:
    for row in range(1, min(ws.max_row, 20) + 1):
        if text(ws.cell(row, 1).value).casefold() != "station":
            continue
        years = {}
        for col in range(2, ws.max_column + 1):
            year = year_from_header(ws.cell(row, col).value)
            if year:
                years[col] = year
        if len(years) >= 5:
            return row, years
    return None, {}


def find_month_header(ws) -> tuple[int, dict[int, str]] | tuple[None, dict]:
    for row in range(1, min(ws.max_row, 20) + 1):
        if text(ws.cell(row, 1).value).casefold() not in {"monitoring station", "station"}:
            continue
        months = {}
        for col in range(2, ws.max_column + 1):
            value = text(ws.cell(row, col).value)
            if value in MONTHS:
                months[col] = value
        if len(months) >= 10:
            return row, months
    return None, {}


def sheet_notes(ws, start_row: int) -> tuple[list[str], str, str]:
    notes = []
    data_as_of = ""
    source = ""
    in_notes = False
    for row in range(start_row, ws.max_row + 1):
        value = text(ws.cell(row, 1).value)
        if not value:
            continue
        low = value.casefold()
        if low.startswith("notes"):
            in_notes = True
            remainder = value[5:].strip(" :")
            if remainder:
                notes.append(remainder)
            continue
        if low.startswith("data as of"):
            data_as_of = re.sub(r"^data as of\s*:?\s*", "", value, flags=re.IGNORECASE)
            in_notes = False
            continue
        if low.startswith("source"):
            source = re.sub(r"^source\s*:?\s*", "", value, flags=re.IGNORECASE)
            continue
        if in_notes:
            notes.append(value)
    return notes, data_as_of, source


def infer_zone(station: str) -> str:
    low = station.casefold()
    for zone in ("West Bay", "East Bay", "Central Bay", "South Bay"):
        if zone.casefold() in low:
            return zone
    return ""


def parse_water_workbook(key: str) -> tuple[list[dict], list[dict], list[dict], dict]:
    spec = VARIABLES[key]
    path = find_workbook(spec["workbook_key"])
    workbook = openpyxl.load_workbook(path, data_only=True)
    long_rows = []
    station_rows = []
    source_meta = None
    for ws in workbook.worksheets:
        header_row, year_cols = find_water_header(ws)
        if header_row is None:
            workbook.close()
            raise RuntimeError(f"Could not find year header in {path.name} / {ws.title}")
        notes, data_as_of, source = sheet_notes(ws, header_row + 1)
        source_meta = {
            "file": path.name,
            "sheet": ws.title,
            "data_as_of": data_as_of,
            "source": source,
            "notes": notes,
        }
        group = None
        subregion = ""
        for row in range(header_row + 1, ws.max_row + 1):
            name = text(ws.cell(row, 1).value)
            if not name:
                continue
            low = name.casefold()
            if low.startswith(("notes", "source", "data as of")):
                break
            if low == "laguna de bay":
                group = "lake"
                subregion = ""
                continue
            if low == "major tributaries":
                group = "tributary"
                subregion = ""
                continue
            if low in {"ncr", "laguna", "rizal"} or low.startswith("region "):
                subregion = name
                continue
            if group is None:
                group = "unclassified"
            station_type = group
            zone = infer_zone(name) if station_type == "lake" else ""
            station_rows.append({
                "station_name": name,
                "station_type": station_type,
                "subregion": subregion,
                "reported_lake_zone": zone,
                "variable": spec["name"],
                "source_id": WQ_SOURCE_IDS[key],
                "first_source_year": min(year_cols.values()),
                "last_source_year": max(year_cols.values()),
            })
            for col, year in sorted(year_cols.items(), key=lambda item: item[1]):
                raw = ws.cell(row, col).value
                numeric = is_number(raw)
                if numeric:
                    value = float(raw)
                    quality = "ok"
                    if not (spec["low"] <= value <= spec["high"]):
                        quality = "out_of_range"
                else:
                    value = ""
                    quality = QUALITY_LABELS.get(missing_kind(raw), "unexpected_text")
                long_rows.append({
                    "station_name": name,
                    "station_type": station_type,
                    "subregion": subregion,
                    "reported_lake_zone": zone,
                    "year": year,
                    "variable": spec["name"],
                    "variable_key": key,
                    "value": value,
                    "unit": spec["unit"],
                    "quality_flag": quality,
                    "source_id": WQ_SOURCE_IDS[key],
                    "source_file": path.name,
                    "source_sheet": ws.title,
                    "source_row": row,
                    "source_data_as_of": data_as_of,
                })
    workbook.close()
    return long_rows, station_rows, notes, source_meta


def build_wide(long_rows: list[dict], station_types: set[str] | None = None) -> list[dict]:
    by_station_year = {}
    station_meta = {}
    for row in long_rows:
        if station_types and row["station_type"] not in station_types:
            continue
        key = (row["station_name"], row["year"])
        if key not in by_station_year:
            by_station_year[key] = {}
            station_meta[key] = row
        if row["quality_flag"] in {"ok", "out_of_range"} and row["value"] != "":
            by_station_year[key][row["variable"]] = row["value"]
            by_station_year[key][f"{row['variable']}__flag"] = row["quality_flag"]
        else:
            by_station_year[key][f"{row['variable']}__flag"] = row["quality_flag"]

    wide = []
    for (station, year), values in sorted(by_station_year.items(), key=lambda item: (item[0][0], item[0][1])):
        meta = station_meta[(station, year)]
        record = {
            "station_name": station,
            "station_type": meta["station_type"],
            "subregion": meta["subregion"],
            "reported_lake_zone": meta["reported_lake_zone"],
            "year": year,
        }
        numeric_count = 0
        for key, spec in VARIABLES.items():
            variable = spec["name"]
            value = values.get(variable, "")
            record[variable] = value
            record[f"{variable}__flag"] = values.get(f"{variable}__flag", "not_in_source")
            if value != "":
                numeric_count += 1
        record["variables_reported"] = numeric_count
        record["all_six_reported"] = "yes" if numeric_count == 6 else "no"
        record["source_data_as_of"] = meta["source_data_as_of"]
        wide.append(record)
    return wide


def build_station_registry(station_rows: list[dict], long_rows: list[dict]) -> list[dict]:
    numeric_years = defaultdict(set)
    for row in long_rows:
        if row["quality_flag"] in {"ok", "out_of_range"} and row["value"] != "":
            numeric_years[row["station_name"]].add(row["year"])
    grouped = {}
    for row in station_rows:
        name = row["station_name"]
        if name not in grouped:
            grouped[name] = {
                "station_name": name,
                "station_type": row["station_type"],
                "subregion": row["subregion"],
                "reported_lake_zone": row["reported_lake_zone"],
                "first_source_year": row["first_source_year"],
                "last_source_year": row["last_source_year"],
                "variables_present_in_source": [],
            }
        grouped[name]["variables_present_in_source"].append(row["variable"])
        grouped[name]["reported_lake_zone"] = grouped[name]["reported_lake_zone"] or row["reported_lake_zone"]
    rows = []
    for name, row in sorted(grouped.items(), key=lambda item: (item[1]["station_type"], item[0])):
        variables = sorted(set(row.pop("variables_present_in_source")))
        years = sorted(numeric_years.get(name, set()))
        row["variables_present_in_source"] = "; ".join(variables)
        row["n_variables_in_source"] = len(variables)
        row["first_numeric_year"] = years[0] if years else ""
        row["last_numeric_year"] = years[-1] if years else ""
        row["numeric_years_count"] = len(years)
        row["zone_evidence"] = (
            "Lake-zone name stated in the station label" if row["reported_lake_zone"] else
            "No lake-zone name in source; no coordinates supplied"
        )
        rows.append(row)
    return rows


def coverage_summary(long_rows: list[dict]) -> list[dict]:
    groups = defaultdict(lambda: {"cells": 0, "numeric": 0, "missing": 0, "out_of_range": 0})
    for row in long_rows:
        key = (row["variable"], row["station_type"])
        groups[key]["cells"] += 1
        if row["quality_flag"] in {"ok", "out_of_range"} and row["value"] != "":
            groups[key]["numeric"] += 1
            if row["quality_flag"] == "out_of_range":
                groups[key]["out_of_range"] += 1
        else:
            groups[key]["missing"] += 1
    rows = []
    for (variable, station_type), counts in sorted(groups.items()):
        rows.append({
            "variable": variable,
            "station_type": station_type,
            "station_year_cells": counts["cells"],
            "numeric_cells": counts["numeric"],
            "missing_cells": counts["missing"],
            "missing_percent": round(100 * counts["missing"] / counts["cells"], 2),
            "out_of_range_cells": counts["out_of_range"],
        })
    return rows


def lake_coverage(wide: list[dict]) -> list[dict]:
    rows = []
    for row in wide:
        if row["station_type"] != "lake":
            continue
        rows.append({
            "station_name": row["station_name"],
            "reported_lake_zone": row["reported_lake_zone"],
            "year": row["year"],
            "nitrate": "yes" if row["Nitrate"] != "" else "no",
            "phosphate": "yes" if row["Phosphate"] != "" else "no",
            "bod": "yes" if row["BOD"] != "" else "no",
            "ph": "yes" if row["pH"] != "" else "no",
            "temperature": "yes" if row["Water temperature"] != "" else "no",
            "do": "yes" if row["Dissolved oxygen (DO)"] != "" else "no",
            "variables_reported": row["variables_reported"],
            "complete_six": row["all_six_reported"],
        })
    return rows


def parse_weather_workbook(key: str) -> tuple[list[dict], list[dict], list[dict]]:
    spec = WEATHER_SPECS[key]
    path = find_workbook(spec["workbook_key"])
    workbook = openpyxl.load_workbook(path, data_only=True)
    long_rows = []
    station_rows = []
    coverage_rows = []
    for ws in workbook.worksheets:
        header_row, month_cols = find_month_header(ws)
        if header_row is None:
            workbook.close()
            raise RuntimeError(f"Could not find month header in {path.name} / {ws.title}")
        year = int(ws.title)
        notes, data_as_of, source = sheet_notes(ws, header_row + 1)
        provisional = "yes" if year == 2025 else "no"
        numeric_by_station = Counter()
        possible_by_station = Counter()
        for row in range(header_row + 1, ws.max_row + 1):
            name = text(ws.cell(row, 1).value)
            if not name:
                continue
            if name.casefold().startswith(("notes", "source", "data as of")):
                break
            station_rows.append({
                "station_name": name,
                "variable": spec["name"],
                "year": year,
                "source_id": WEATHER_SOURCE_IDS[key],
            })
            for col, month_name in sorted(month_cols.items(), key=lambda item: MONTH_NUMBER[item[1]]):
                raw = ws.cell(row, col).value
                month_number = MONTH_NUMBER[month_name]
                numeric = is_number(raw)
                possible_by_station[name] += 1
                if numeric:
                    value = float(raw)
                    quality = "ok"
                    if not (0 <= value <= (5000 if key == "rainfall" else 60)):
                        quality = "out_of_range"
                    numeric_by_station[name] += 1
                else:
                    value = ""
                    quality = QUALITY_LABELS.get(missing_kind(raw), "unexpected_text")
                long_rows.append({
                    "station_name": name,
                    "year": year,
                    "month": month_name,
                    "month_number": month_number,
                    "period_start": f"{year}-{month_number:02d}-01",
                    "variable": spec["name"],
                    "variable_key": key,
                    "value": value,
                    "unit": spec["unit"],
                    "quality_flag": quality,
                    "provisional_2025": provisional,
                    "source_id": WEATHER_SOURCE_IDS[key],
                    "source_file": path.name,
                    "source_sheet": ws.title,
                    "source_row": row,
                    "source_data_as_of": data_as_of,
                })
        for station, possible in possible_by_station.items():
            numeric = numeric_by_station[station]
            coverage_rows.append({
                "variable": spec["name"],
                "year": year,
                "station_name": station,
                "months_in_source_sheet": possible,
                "numeric_months": numeric,
                "missing_months": possible - numeric,
                "complete_12_months": "yes" if numeric == 12 else "no",
                "provisional_2025": provisional,
            })
    workbook.close()
    return long_rows, station_rows, coverage_rows


def build_weather_wide(rain_rows: list[dict], wind_rows: list[dict]) -> list[dict]:
    merged = {}
    for variable_key, rows, column in (
        ("rainfall", rain_rows, "rainfall_mm"),
        ("wind", wind_rows, "wind_speed_ms"),
    ):
        for row in rows:
            key = (row["station_name"], row["year"], row["month_number"])
            if key not in merged:
                merged[key] = {
                    "station_name": row["station_name"],
                    "year": row["year"],
                    "month": row["month"],
                    "month_number": row["month_number"],
                    "period_start": row["period_start"],
                    "rainfall_mm": "",
                    "wind_speed_ms": "",
                    "rainfall_quality": "",
                    "wind_quality": "",
                    "provisional_2025": row["provisional_2025"],
                }
            if row["quality_flag"] in {"ok", "out_of_range"} and row["value"] != "":
                merged[key][column] = row["value"]
            merged[key][f"{variable_key}_quality"] = row["quality_flag"]
    return [merged[key] for key in sorted(merged, key=lambda k: (k[0], k[1], k[2]))]


def write_cleaning_report(
    long_rows: list[dict],
    lake_rows: list[dict],
    weather_long: dict[str, list[dict]],
    water_cov: list[dict],
    weather_cov: list[dict],
) -> None:
    lake_years = len(lake_rows)
    complete = sum(1 for row in lake_rows if row["all_six_reported"] == "yes")
    lines = [
        "# Environmental data cleaning report",
        "",
        "Generated by `python/clean_data.py`. Source workbooks were read only.",
        "",
        "## Water quality",
        "",
        f"- Long-format rows: **{len(long_rows):,}** (one row per station-year-variable, including missing cells).",
        f"- Lake station-years: **{lake_years}**; complete across all six variables: **{complete}**.",
        "- Values are annual station summaries. The aggregation method is not stated in the source.",
        "- Missing markers are retained as quality flags, never converted to zero.",
        "- `**` is treated as station no longer monitoring; the U+2026 no-data symbol as no data; `...` as undefined missing.",
        "- Lake-zone labels are copied only when they appear in the station name. Ten stations have no",
        "  zone label in the source and no coordinates, so their zone is left blank.",
        "- There are 59 distinct station labels (15 lake, 44 tributary). The 15 lake station names are",
        "  identical across all six workbooks. On the tributary side, pH and temperature omit four sites",
        "  and call 'Angono River' what the nutrient/BOD files call 'Angono'. That name pair is not",
        "  merged automatically; a documented crosswalk is required.",
        "- 2025 values are present, but the source does not document how many months contributed to each",
        "  2025 annual value.",
        "",
        "### Water-quality coverage",
        "",
        "| Variable | Station type | Cells | Numeric | Missing | Missing % | Out of range |",
        "|---|---|---:|---:|---:|---:|---:|",
    ]
    for row in water_cov:
        lines.append(
            f"| {row['variable']} | {row['station_type']} | {row['station_year_cells']:,} | "
            f"{row['numeric_cells']:,} | {row['missing_cells']:,} | {row['missing_percent']} | "
            f"{row['out_of_range_cells']} |"
        )
    lines += [
        "",
        "## Weather",
        "",
        "| Variable | Numeric station-month cells | Missing | Missing % | 2025 note |",
        "|---|---:|---:|---:|---|",
    ]
    for key in ("rainfall", "wind"):
        rows = weather_long[key]
        numeric = sum(1 for row in rows if row["value"] != "")
        missing = len(rows) - numeric
        pct = round(100 * missing / len(rows), 2) if rows else ""
        lines.append(
            f"| {WEATHER_SPECS[key]['name']} | {numeric:,} | {missing:,} | {pct} | "
            "Partial (most stations Jan-Jun; some 12 months); source says 2025 data are still being verified |"
        )
    lines += [
        "",
        "- `period_start` is the first day of the summary month, not a measurement date.",
        "- Rainfall and wind station rosters are not identical across years; the wide file preserves",
        "  the union rather than dropping unmatched stations.",
        "- No PAGASA station has coordinates in the source files and there is no documented mapping",
        "  from weather stations to lake zones.",
        "",
        "## Prohibited joins",
        "",
        "No row in this output should be joined to an annual water-quality value as if it were a",
        "same-day observation. Monthly weather and annual water quality are different temporal",
        "supports. Any future analysis must state the aggregation and lag explicitly.",
    ]
    (TABLE_DIR / "cleaning_report.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    ensure_dirs()
    all_long = []
    all_stations = []
    for key in VARIABLES:
        long_rows, station_rows, _notes, _meta = parse_water_workbook(key)
        all_long.extend(long_rows)
        all_stations.extend(station_rows)

    wide_all = build_wide(all_long)
    wide_lake = build_wide(all_long, {"lake"})
    station_registry = build_station_registry(all_stations, all_long)
    water_cov = coverage_summary(all_long)
    lake_cov = lake_coverage(wide_all)

    write_csv(PROCESSED_DIR / "water_quality_long.csv", all_long)
    write_csv(PROCESSED_DIR / "water_quality_all_annual_wide.csv", wide_all)
    write_csv(PROCESSED_DIR / "water_quality_lake_annual.csv", wide_lake)
    write_csv(PROCESSED_DIR / "station_registry.csv", station_registry)
    write_csv(PROCESSED_DIR / "water_quality_coverage_summary.csv", water_cov)
    write_csv(PROCESSED_DIR / "lake_station_year_coverage.csv", lake_cov)

    weather_long = {}
    weather_station_rows = []
    weather_cov = []
    for key in WEATHER_SPECS:
        rows, stations, coverage = parse_weather_workbook(key)
        weather_long[key] = rows
        weather_station_rows.extend(stations)
        weather_cov.extend(coverage)
        write_csv(PROCESSED_DIR / f"{key}_monthly.csv", rows)

    weather_wide = build_weather_wide(weather_long["rainfall"], weather_long["wind"])
    write_csv(PROCESSED_DIR / "weather_monthly_wide.csv", weather_wide)
    write_csv(PROCESSED_DIR / "weather_station_registry.csv", weather_station_rows)
    write_csv(PROCESSED_DIR / "weather_coverage_summary.csv", weather_cov)

    write_cleaning_report(all_long, wide_lake, weather_long, water_cov, weather_cov)

    print("Cleaning complete")
    print(f"Water-quality long rows : {len(all_long):,}")
    print(f"Wide water station-years: {len(wide_all):,}")
    print(f"Lake station-years       : {len(wide_lake):,}")
    print(f"Lake complete six vars   : {sum(1 for row in wide_lake if row['all_six_reported'] == 'yes')}")
    print(f"Rainfall monthly rows    : {len(weather_long['rainfall']):,}")
    print(f"Wind monthly rows        : {len(weather_long['wind']):,}")
    print(f"Weather wide rows        : {len(weather_wide):,}")


if __name__ == "__main__":
    main()
