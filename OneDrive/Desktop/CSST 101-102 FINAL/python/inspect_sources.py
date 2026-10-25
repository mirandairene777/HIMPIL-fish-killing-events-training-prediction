"""
HIMPIL Phase 1 - inspect and document every supplied source.

This script reads the eight Excel workbooks, two PDFs, two Word documents,
the uploaded React UI archive, and the archived 2019 UPD article citation.
It does not alter any source file.

Outputs
-------
reports/tables/source_inventory.csv
reports/tables/source_manifest.md
data/processed/source_text/*.txt        (searchable text extracts)
data/processed/source_text/zip_manifest.txt

Run from the project root:
    .venv\\Scripts\\python.exe python\\inspect_sources.py
"""

from __future__ import annotations

import hashlib
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from common import (
    DOC_TEXT_DIR,
    PROJECT_ROOT,
    RAW_DIR,
    SOURCE_DOC_DIR,
    TABLE_DIR,
    ensure_dirs,
    markdown_table,
    read_csv,
    text,
    write_csv,
)

SOURCE_COLUMNS = [
    "source_id",
    "filename",
    "relative_path",
    "source_type",
    "source_organization",
    "publication_date",
    "reporting_period",
    "geography",
    "content_type",
    "individual_incident_records",
    "incident_dates",
    "incident_locations",
    "environmental_measurements",
    "units_documented",
    "causes_or_losses",
    "fish_kill_label_suitability",
    "pages_or_sheets",
    "used_in_himpil_pipeline",
    "sha256",
    "notes",
]

# Fixed source IDs make later audit tables stable.
SOURCE_IDS = {
    "Table 1.2 Amount of Rainfall": "S01",
    "Table 1.32 Concentration Level of Nitrate": "S02",
    "Table 1.33 Concentration Level of Phosphate": "S03",
    "Table 1.34.2 Concentration Level of BOD": "S04",
    "Table 1.37 Levels of pH": "S05",
    "Table 1.38 Temperature": "S06",
    "Table 1.39.2 Concentration Level of DO": "S07",
    "Table 1.6 Wind Speed": "S08",
    "AWQR_2009-2012.pdf": "S09",
    "sustainability-16-10521.pdf": "S10",
    "PROPOSAL_-HIMPIL-1.docx": "S11",
    "HIMPIL.docx": "S12",
    "himpil-ui.zip": "S13",
    "upd_fish_kill_2019_article.md": "S14",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_id(path: Path) -> str:
    for needle, value in SOURCE_IDS.items():
        if needle in path.name:
            return value
    raise KeyError(f"No source ID assigned for {path.name}")


def workbook_metadata(path: Path, inventory: list[dict]) -> dict:
    rows = [row for row in inventory if row["file"] == path.name]
    if not rows:
        raise RuntimeError(f"{path.name} is missing from dataset_inventory.csv; run inspect_data.py")
    variables = sorted({row["variable"] for row in rows})
    stations = [int(row["station_count"]) for row in rows if row["station_count"]]
    cells = sum(int(row["total_station_period_cells"]) for row in rows)
    numeric = sum(int(row["numeric_cells"]) for row in rows)
    missing_pct = round(100 * (cells - numeric) / cells, 2) if cells else ""
    frequencies = sorted({row["frequency"] for row in rows})
    notes = sorted({row["notes"] for row in rows if row["notes"]})
    if "Rainfall" in variables:
        return {
            "source_organization": "Philippine Atmospheric, Geophysical and Astronomical Services Administration (PAGASA), DOST",
            "publication_date": "Data as of 2026-07-09 (source note)",
            "reporting_period": "2016-2025; 2025 incomplete and still undergoing verification",
            "geography": "PAGASA synoptic monitoring stations nationwide",
            "content_type": "Aggregated monthly station totals (one value per station-month)",
            "individual_incident_records": "None",
            "incident_dates": "",
            "incident_locations": "",
            "environmental_measurements": "Rainfall",
            "units_documented": "Yes - millimetres (source title and sheet note)",
            "causes_or_losses": "",
            "fish_kill_label_suitability": "Not suitable - no fish kill records; monthly summaries only",
            "pages_or_sheets": f"{len(rows)} annual worksheets",
            "used_in_himpil_pipeline": "Yes - weather context/descriptive analysis only",
            "notes": (
                f"{min(stations)}-{max(stations)} stations per year; {numeric:,}/{cells:,} numeric "
                f"station-month cells ({missing_pct}% missing). 2025 is provisional and incomplete: "
                "most stations have January-June values only, while some have all 12 months; the "
                "source says remaining synoptic data are still undergoing verification. "
                "No station coordinates are supplied."
            ),
        }
    if "Wind speed" in variables:
        return {
            "source_organization": "Philippine Atmospheric, Geophysical and Astronomical Services Administration (PAGASA), DOST",
            "publication_date": "Data as of 2026-07-09 (source note)",
            "reporting_period": "2016-2025; 2025 incomplete and still undergoing verification",
            "geography": "PAGASA synoptic monitoring stations nationwide",
            "content_type": "Aggregated monthly station values (one value per station-month)",
            "individual_incident_records": "None",
            "incident_dates": "",
            "incident_locations": "",
            "environmental_measurements": "Wind speed (average over a 10-minute period)",
            "units_documented": "Yes - metres per second (source title and sheet note)",
            "causes_or_losses": "",
            "fish_kill_label_suitability": "Not suitable - no fish kill records; monthly summaries only",
            "pages_or_sheets": f"{len(rows)} annual worksheets",
            "used_in_himpil_pipeline": "Yes - weather context/descriptive analysis only",
            "notes": (
                f"{min(stations)}-{max(stations)} stations per year; {numeric:,}/{cells:,} numeric "
                f"station-month cells ({missing_pct}% missing). 2025 is provisional and incomplete: "
                "most stations have January-June values only, while some have all 12 months; the "
                "source says remaining synoptic data are still undergoing verification. "
                "No station coordinates are supplied."
            ),
        }
    return {
        "source_organization": "Laguna Lake Development Authority (LLDA), DENR",
        "publication_date": "Data as of 2026-07-28 (source note)",
        "reporting_period": "2016-2025",
        "geography": "Laguna de Bay and its major tributaries",
        "content_type": "Aggregated annual station summaries (one value per station-year)",
        "individual_incident_records": "None",
        "incident_dates": "",
        "incident_locations": "",
        "environmental_measurements": ", ".join(variables),
        "units_documented": "Yes - see the workbook title/unit line",
        "causes_or_losses": "",
        "fish_kill_label_suitability": "Not suitable - no fish kill records; annual aggregation",
        "pages_or_sheets": f"{len(rows)} worksheet(s)",
        "used_in_himpil_pipeline": "Yes - cleaned annual environmental data",
        "notes": (
            f"{min(stations)}-{max(stations)} stations per year; {numeric:,}/{cells:,} numeric "
            f"station-year cells ({missing_pct}% missing). Aggregation method (mean/maximum/other) "
            "is not stated. No observation dates or coordinates."
            + (" Workbook notes: " + " | ".join(notes) if notes else "")
        ),
    }


def pdf_metadata(path: Path) -> tuple[dict, str]:
    reader = PdfReader(str(path))
    meta = reader.metadata or {}
    parts = []
    for index, page in enumerate(reader.pages, start=1):
        parts.append(f"\n\n===== PAGE {index} =====\n\n{page.extract_text() or ''}")
    extracted = "".join(parts)
    dates = {key: text(value) for key, value in meta.items() if key in {"/CreationDate", "/ModDate"}}
    return {
        "pages": len(reader.pages),
        "metadata": {text(k): text(v) for k, v in meta.items()},
        "metadata_dates": dates,
        "text": extracted,
    }, extracted


def docx_text(path: Path) -> str:
    document = Document(path)
    parts = ["PARAGRAPHS"]
    parts.extend(paragraph.text for paragraph in document.paragraphs)
    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"\nTABLE {table_index}")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def build_source_rows(inventory: list[dict]) -> list[dict]:
    paths = []
    paths.extend(sorted(RAW_DIR.glob("*.xlsx")))
    paths.extend(sorted(SOURCE_DOC_DIR.glob("*")))
    rows = []
    zip_manifest_lines = []
    for path in paths:
        if not path.is_file() or path.name.startswith("."):
            continue
        sid = source_id(path)
        row = {
            "source_id": sid,
            "filename": path.name,
            "relative_path": str(path.relative_to(PROJECT_ROOT)),
            "source_type": path.suffix.lower().lstrip("."),
            "sha256": sha256(path),
        }
        if path.suffix.lower() == ".xlsx":
            row.update(workbook_metadata(path, inventory))
        elif path.suffix.lower() == ".pdf":
            info, extracted = pdf_metadata(path)
            (DOC_TEXT_DIR / f"{path.stem}.txt").write_text(extracted, encoding="utf-8")
            row["pages_or_sheets"] = f"{info['pages']} pages"
            row["sha256"] = sha256(path)
            if "AWQR" in path.name:
                row.update({
                    "source_organization": "Laguna Lake Development Authority (LLDA), DENR",
                    "publication_date": "Report covers 2009-2012; PDF metadata created 2013-11-29",
                    "reporting_period": "2009-2012",
                    "geography": "Laguna de Bay and its tributary rivers",
                    "content_type": "Annual water-quality report with narrative, annual summary tables and graphs",
                    "individual_incident_records": "None",
                    "incident_dates": "",
                    "incident_locations": "",
                    "environmental_measurements": (
                        "17 physico-chemical parameters, 3 biological parameters, 2 microbiological "
                        "parameters and lake primary productivity; includes DO, temperature, pH, nitrate, "
                        "phosphate and BOD"
                    ),
                    "units_documented": "Yes",
                    "causes_or_losses": "",
                    "fish_kill_label_suitability": "Not suitable - no fish kill incident records; annual summaries",
                    "used_in_himpil_pipeline": "Background/provenance only; not merged",
                    "notes": (
                        "Foreword says lake and tributary stations were sampled once a month, but the "
                        "published tables and graphs present annual averages by station, not dated samples. "
                        "Only five lake stations existed in 2009-2011; four more were added in 2012. "
                        "The report contains no individual fish kill event records. Station maps are images; "
                        "no station coordinates are available in machine-readable form."
                    ),
                })
            else:
                row.update({
                    "source_organization": "Sustainability (MDPI); authors from UP Diliman, UST, and regional partners",
                    "publication_date": "2024-11-30",
                    "reporting_period": "Literature review covering events from 1972 to 2023",
                    "geography": "Southeast Asia, with country-level tables for the Philippines",
                    "content_type": "Peer-reviewed narrative review with event tables",
                    "individual_incident_records": (
                        "None usable for Laguna de Bay: one aggregate 1972-1998 statement and one "
                        "contradictory 2021 table row"
                    ),
                    "incident_dates": (
                        "Laguna de Bay: '1972-1998, May to September' (aggregate); "
                        "2021 January row is contradicted by its cited article"
                    ),
                    "incident_locations": (
                        "Laguna de Bay (aggregate); central arm and west arm mentioned for the aggregate period. "
                        "The 2021 row's citation refers to Sampaloc Lake, not Laguna de Bay."
                    ),
                    "environmental_measurements": "None - review article",
                    "units_documented": "Not applicable",
                    "causes_or_losses": (
                        "Aggregate Laguna statement: low dissolved oxygen; over half of those cases associated "
                        "with blue-green phytoplankton blooms; other causes reported as agricultural/industrial "
                        "pollution, fish pathogens and other causes. No event-level losses for Laguna de Bay."
                    ),
                    "fish_kill_label_suitability": (
                        "Not suitable - the Laguna evidence is aggregate; the 2021 row conflicts with its source"
                    ),
                    "used_in_himpil_pipeline": "Yes - incident evidence audit only",
                    "notes": (
                        "Section 2.2 and Table 3 were checked. Table 3 row '2021 Jan. Laguna de Bay / Tilapia / "
                        "Low DO levels [121]' cites reference 121, whose title is '2 Tons of Tilapia Lost as Fish "
                        "Kill Hits Laguna's Sampaloc Lake'. Sampaloc Lake is a different waterbody. The row was "
                        "therefore not accepted as a Laguna de Bay event. DOI: 10.3390/su162310521."
                    ),
                })
        elif path.suffix.lower() == ".docx":
            extracted = docx_text(path)
            (DOC_TEXT_DIR / f"{path.stem}.txt").write_text(extracted, encoding="utf-8")
            document = Document(path)
            row["pages_or_sheets"] = f"{len(document.paragraphs)} paragraphs; {len(document.tables)} tables"
            if "PROPOSAL" in path.name:
                row.update({
                    "source_organization": "HIMPIL student research team",
                    "publication_date": "Undated; DOCX metadata modified 2026-10-19",
                    "reporting_period": "Not applicable (research proposal)",
                    "geography": "Laguna de Bay, Philippines",
                    "content_type": "Narrative research proposal / concept paper",
                    "individual_incident_records": "None",
                    "incident_dates": "",
                    "incident_locations": "",
                    "environmental_measurements": "Proposed inputs only: air temperature, rainfall, wind speed/direction, cloud cover, solar radiation, water temperature, DO",
                    "units_documented": "No data tables; some units named in narrative",
                    "causes_or_losses": "Describes the hypothesised stratification/oxygen-depletion mechanism; no event records or losses",
                    "fish_kill_label_suitability": "Not suitable - proposal contains no observations or labels",
                    "used_in_himpil_pipeline": "Requirements/background only",
                    "notes": "Proposes a 24-72 hour forecasting system but supplies no dated observations or event data.",
                })
            else:
                row.update({
                    "source_organization": "Laguna State Polytechnic University student output",
                    "publication_date": "Undated; AY 2026-2027 stated in document",
                    "reporting_period": "Not applicable",
                    "geography": "Laguna de Bay",
                    "content_type": "Research topic discovery form with 9 tables",
                    "individual_incident_records": "None",
                    "incident_dates": "",
                    "incident_locations": "",
                    "environmental_measurements": "None",
                    "units_documented": "No",
                    "causes_or_losses": "Problem description only",
                    "fish_kill_label_suitability": "Not suitable - no data",
                    "used_in_himpil_pipeline": "Background only",
                    "notes": "Concept/ideation document; no environmental observations or fish kill records.",
                })
        elif path.suffix.lower() == ".zip":
            with zipfile.ZipFile(path) as archive:
                names = sorted(archive.namelist())
            zip_manifest_lines = names
            (DOC_TEXT_DIR / "zip_manifest.txt").write_text("\n".join(names), encoding="utf-8")
            row.update({
                "source_organization": "Uploaded HIMPIL UI prototype (React/Vite)",
                "publication_date": "Not stated",
                "reporting_period": "Not applicable",
                "geography": "Laguna de Bay (interface labels only)",
                "content_type": "Front-end prototype; static placeholder screens",
                "individual_incident_records": "None",
                "incident_dates": "",
                "incident_locations": "",
                "environmental_measurements": "None - all metric values are placeholders",
                "units_documented": "Labels only (mg/L, deg C, mm, m/s)",
                "causes_or_losses": "",
                "fish_kill_label_suitability": "Not applicable - no data",
                "pages_or_sheets": f"{len(names)} archived files",
                "used_in_himpil_pipeline": "Yes - dashboard shell only",
                "notes": "No backend, API, dataset, coordinates, or model output. Placeholder text says data are not connected.",
            })
        elif path.suffix.lower() == ".md":
            row.update({
                "source_organization": "University of the Philippines Diliman Information Office",
                "publication_date": "2019-06-10 (page also shows a June 13 dateline)",
                "reporting_period": "2019-05-04 to about 2019-05-31",
                "geography": "Biñan and Pila, Laguna, Laguna de Bay",
                "content_type": "University news report with an LLDA interview and a water-quality results table",
                "individual_incident_records": "1 event report across two municipalities (Biñan and Pila)",
                "incident_dates": "Biñan start 2019-05-04; Pila start 2019-05-12; reported to last until end-May 2019",
                "incident_locations": "Biñan, Laguna; Barangay Linga, Pila, Laguna",
                "environmental_measurements": "Biñan water sample on 2019-05-29: DO 4.7 mg/L; ammonia 0.16 mg/L; nitrate 0.4 mg/L; phosphate 1.00 mg/L",
                "units_documented": "Yes - mg/L",
                "causes_or_losses": (
                    "LLDA specialist attributed the kill to a green tide/algal bloom and oxygen deficiency. "
                    "Estimated 500,000 dead fish in Biñan cages; 60% tilapia and 80% bangus mortality reported "
                    "in Pila. No peso loss given."
                ),
                "fish_kill_label_suitability": "Partial only - one dated event/location; no dated pre-event environmental series",
                "pages_or_sheets": "Archived web citation (Markdown snapshot)",
                "used_in_himpil_pipeline": "Yes - event extraction and matching audit",
                "notes": (
                    "The 29 May 2019 sample was collected during/at the end of the event, not before it. "
                    "Article is secondary reporting with agency attribution, not a formal investigation report. "
                    "URL: https://upd.edu.ph/seasonal-fish-kill-incident-leaves-thousands-of-unprofitable-commodity-fish-in-laguna-de-bay/"
                ),
            })
        rows.append(row)
    return rows


def write_manifest(rows: list[dict]) -> None:
    columns = [
        ("source_id", "ID"),
        ("filename", "Source"),
        ("source_type", "Type"),
        ("publication_date", "Publication / data date"),
        ("reporting_period", "Reporting period"),
        ("content_type", "Content"),
        ("individual_incident_records", "Individual incidents?"),
        ("fish_kill_label_suitability", "Usable labels?"),
    ]
    table = markdown_table(rows, columns)
    detail_lines = []
    for row in rows:
        detail_lines += [
            f"### {row['source_id']} - {row['filename']}",
            "",
            f"- **Organization:** {row['source_organization']}",
            f"- **Publication/data date:** {row['publication_date']}",
            f"- **Reporting period:** {row['reporting_period']}",
            f"- **Geography:** {row['geography']}",
            f"- **Content:** {row['content_type']}",
            f"- **Incidents:** {row['individual_incident_records']}",
            f"- **Incident dates:** {row['incident_dates'] or 'Not applicable'}",
            f"- **Incident locations:** {row['incident_locations'] or 'Not applicable'}",
            f"- **Environmental measurements:** {row['environmental_measurements'] or 'None'}",
            f"- **Units:** {row['units_documented']}",
            f"- **Causes/losses:** {row['causes_or_losses'] or 'Not applicable'}",
            f"- **Label suitability:** {row['fish_kill_label_suitability']}",
            f"- **Used in pipeline:** {row['used_in_himpil_pipeline']}",
            f"- **Notes:** {row['notes'] or 'None'}",
            f"- **SHA-256:** `{row['sha256']}`",
            "",
        ]
    manifest = "\n".join([
        "# HIMPIL source inventory and manifest",
        "",
        "Generated by `python/inspect_sources.py`. Original source files are unchanged.",
        "",
        "## Inventory",
        "",
        table,
        "",
        "## Critical incident-evidence finding",
        "",
        "The uploaded review contains only an **aggregate** Laguna de Bay statement for 1972-1998 "
        "(May-September; low dissolved oxygen and blue-green algal blooms). An aggregate statement "
        "cannot be expanded into individual events. Its one 2021 'Laguna de Bay' table row cites a "
        "news article about **Sampaloc Lake**, a different waterbody, so that row was not accepted. "
        "The added UPD article supplies one dated event report for Biñan and Pila in May 2019. That is "
        "still only one event period and it has no pre-event daily environmental series.",
        "",
        "## Source details",
        "",
        *detail_lines,
    ])
    (TABLE_DIR / "source_manifest.md").write_text(manifest, encoding="utf-8")


def main() -> None:
    ensure_dirs()
    inventory_path = TABLE_DIR / "dataset_inventory.csv"
    if not inventory_path.exists():
        raise SystemExit("Run python/inspect_data.py first to create dataset_inventory.csv")
    inventory = read_csv(inventory_path)
    rows = build_source_rows(inventory)
    rows.sort(key=lambda row: row["source_id"])
    write_csv(TABLE_DIR / "source_inventory.csv", rows, SOURCE_COLUMNS)
    write_manifest(rows)
    print(f"Sources inspected: {len(rows)}")
    print(f"Inventory: {TABLE_DIR / 'source_inventory.csv'}")
    print(f"Manifest : {TABLE_DIR / 'source_manifest.md'}")
    print(f"Text extracts: {DOC_TEXT_DIR}")
    print(f"Timestamp: {datetime.now(timezone.utc).isoformat(timespec='seconds')}")


if __name__ == "__main__":
    main()
