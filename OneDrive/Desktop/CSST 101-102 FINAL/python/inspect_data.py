"""
HIMPIL Phase 2 - source data inspection.

Purpose
-------
Read every Excel workbook in data/raw WITHOUT modifying it, and produce:
  reports/tables/dataset_inventory.csv   one row per worksheet
  reports/tables/data_quality_report.csv one row per worksheet (missing/out-of-range counts)
  reports/tables/station_inventory.csv   one row per station per worksheet
  reports/tables/dataset_inventory.md    human-readable summary

This script only describes the files. It does not clean, merge, impute, label,
or train anything.

Run from the project root:
    .venv\\Scripts\\python.exe python\\inspect_data.py
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_ROOT / "data" / "raw"
TABLE_DIR = PROJECT_ROOT / "reports" / "tables"

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

SECTION_LABELS = {
    "laguna de bay", "major tributaries", "ncr", "laguna", "rizal",
    "region iv-a", "region iv-a (calabarzon)", "calabarzon",
}

VARIABLE_RULES = [
    # (match text, variable name, canonical unit, plausible min, plausible max)
    ("BOD", "Biochemical oxygen demand (BOD)", "mg/L", 0.0, 1000.0),
    ("DO", "Dissolved oxygen (DO)", "mg/L", 0.0, 25.0),
    ("Nitrate", "Nitrate", "mg/L", 0.0, 100.0),
    ("Phosphate", "Phosphate", "mg/L", 0.0, 100.0),
    ("pH", "pH", "pH units", 0.0, 14.0),
    ("Temperature", "Water temperature", "deg C", 0.0, 50.0),
    ("Rainfall", "Rainfall", "mm", 0.0, 5000.0),
    ("Wind Speed", "Wind speed", "m/s", 0.0, 60.0),
]


def text(value) -> str:
    """Return a trimmed string, treating None as an empty string."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalise_name(value) -> str:
    return re.sub(r"\s+", " ", text(value)).casefold()


def variable_info(filename: str):
    for needle, name, unit, low, high in VARIABLE_RULES:
        if needle.lower() in filename.lower():
            return name, unit, low, high
    return "Unknown", "", None, None


def find_header(ws):
    """Find the row that starts with Station / Monitoring Station and has periods."""
    for row in range(1, min(ws.max_row, 30) + 1):
        first = text(ws.cell(row, 1).value).lower()
        if first not in {"station", "monitoring station"}:
            continue
        year_cols = {}
        month_cols = {}
        for col in range(2, ws.max_column + 1):
            value = ws.cell(row, col).value
            if isinstance(value, (int, float)) and 1900 <= value <= 2100:
                year_cols[col] = int(value)
            elif text(value) in MONTHS:
                month_cols[col] = text(value)
        if year_cols or month_cols:
            return row, year_cols, month_cols
    return None, {}, {}


def find_metadata(ws, header_row):
    """Collect title, unit, notes, source and data-as-of text from the sheet."""
    table_title = ""
    unit = ""
    notes = []
    source = ""
    data_as_of = ""
    in_notes = False

    for row in range(1, ws.max_row + 1):
        values = [text(ws.cell(row, c).value) for c in range(1, min(ws.max_column, 10) + 1)]
        joined = " ".join(v for v in values if v)
        if not joined:
            continue
        if row < header_row and not table_title and row > 1:
            table_title = joined
        if row < header_row and not unit:
            match = re.search(r"\(in\s+(.+?)\)", joined, flags=re.IGNORECASE)
            if match:
                unit = match.group(1).strip()
        low = joined.lower()
        if low.startswith("notes"):
            in_notes = True
            note_text = joined[len("notes"):].strip(" :")
            if note_text:
                notes.append(note_text)
            continue
        if in_notes:
            if low.startswith("data as of"):
                data_as_of = re.sub(r"^data as of\s*:?\s*", "", joined, flags=re.IGNORECASE)
                in_notes = False
            elif low.startswith("source"):
                source = re.sub(r"^source\s*:?\s*", "", joined, flags=re.IGNORECASE)
                in_notes = False
            else:
                notes.append(joined)
        if low.startswith("data as of"):
            data_as_of = re.sub(r"^data as of\s*:?\s*", "", joined, flags=re.IGNORECASE)
        if low.startswith("source"):
            source = re.sub(r"^source\s*:?\s*", "", joined, flags=re.IGNORECASE)

    return table_title, unit, notes, source, data_as_of


def is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def missing_marker(value) -> str:
    """Classify a non-numeric cell without inventing a value."""
    raw = text(value)
    if raw == "":
        return "blank"
    if raw == "...":
        return "ellipsis_undefined"
    if raw == "**":
        return "station_discontinued"
    # The source uses a single ellipsis character (U+2026) for "no data".
    # Some sheets instead use three periods; that symbol has no legend, so it is
    # tracked separately rather than silently treated as zero.
    if raw in {"\u2026", "\ufffd", "—", "–"}:
        return "no_data_symbol"
    return f"other:{raw}"


def inspect_sheet(ws, filename, variable, canonical_unit, low, high):
    header_row, year_cols, month_cols = find_header(ws)
    result = {
        "file": filename,
        "sheet": ws.title,
        "sheet_state": ws.sheet_state,
        "table_title": "",
        "variable": variable,
        "source_unit": "",
        "frequency": "",
        "period_type": "",
        "year_start": "",
        "year_end": "",
        "header_row": header_row or "",
        "station_count": 0,
        "lake_station_count": 0,
        "tributary_station_count": 0,
        "duplicate_station_names": "",
        "total_station_period_cells": 0,
        "numeric_cells": 0,
        "missing_cells": 0,
        "missing_percent": "",
        "ellipsis_undefined_cells": 0,
        "no_data_symbol_cells": 0,
        "station_discontinued_cells": 0,
        "blank_cells": 0,
        "other_non_numeric_cells": 0,
        "out_of_range_cells": 0,
        "data_as_of": "",
        "source": "",
        "notes": "",
    }
    if header_row is None:
        result["notes"] = "No station header row found; inspect manually."
        return result, []

    title, unit, notes, source, data_as_of = find_metadata(ws, header_row)
    periods = year_cols if year_cols else month_cols
    result.update({
        "table_title": title,
        "source_unit": unit,
        "data_as_of": data_as_of,
        "source": source,
        "notes": " | ".join(notes),
    })
    if year_cols:
        years = sorted(year_cols.values())
        result["frequency"] = "annual summary (one value per station-year)"
        result["period_type"] = "year"
        result["year_start"] = years[0]
        result["year_end"] = years[-1]
    elif month_cols:
        result["frequency"] = "monthly summary (one value per station-month)"
        result["period_type"] = "month"
        if re.fullmatch(r"\d{4}", ws.title):
            result["year_start"] = result["year_end"] = int(ws.title)

    station_rows = []
    current_group = "unclassified"
    current_subregion = ""
    for row in range(header_row + 1, ws.max_row + 1):
        name = text(ws.cell(row, 1).value)
        if not name:
            continue
        low_name = name.lower()
        if low_name.startswith(("notes", "source", "data as of")):
            break
        if low_name in SECTION_LABELS or low_name.startswith("region "):
            if low_name == "laguna de bay":
                current_group = "Laguna de Bay (lake)"
            elif low_name == "major tributaries":
                current_group = "Tributary"
            elif low_name in {"laguna", "rizal", "ncr"}:
                current_subregion = name
            continue

        values = [(year_cols.get(col) or month_cols.get(col), ws.cell(row, col).value)
                  for col in periods]
        available = 0
        numeric = 0
        out_of_range = 0
        markers = Counter()
        for period, value in values:
            available += 1
            if is_number(value):
                numeric += 1
                if low is not None and high is not None and not (low <= float(value) <= high):
                    out_of_range += 1
            else:
                markers[missing_marker(value)] += 1

        station_rows.append({
            "file": filename,
            "sheet": ws.title,
            "station": name,
            "station_group": current_group,
            "subregion": current_subregion,
            "periods_available": available,
            "numeric_periods": numeric,
            "missing_periods": available - numeric,
            "first_year": min(year_cols.values()) if year_cols else "",
            "last_year": max(year_cols.values()) if year_cols else "",
        })

        result["station_count"] += 1
        if current_group == "Laguna de Bay (lake)":
            result["lake_station_count"] += 1
        elif current_group == "Tributary":
            result["tributary_station_count"] += 1
        result["total_station_period_cells"] += available
        result["numeric_cells"] += numeric
        result["out_of_range_cells"] += out_of_range
        for marker, count in markers.items():
            if marker == "ellipsis_undefined":
                result["ellipsis_undefined_cells"] += count
            elif marker == "no_data_symbol":
                result["no_data_symbol_cells"] += count
            elif marker == "station_discontinued":
                result["station_discontinued_cells"] += count
            elif marker == "blank":
                result["blank_cells"] += count
            else:
                result["other_non_numeric_cells"] += count

    result["missing_cells"] = result["total_station_period_cells"] - result["numeric_cells"]
    if result["total_station_period_cells"]:
        result["missing_percent"] = round(
            100 * result["missing_cells"] / result["total_station_period_cells"], 2
        )

    names = [normalise_name(row["station"]) for row in station_rows]
    duplicates = sorted({name for name, count in Counter(names).items() if count > 1})
    result["duplicate_station_names"] = " | ".join(duplicates)

    if not result["source_unit"]:
        result["source_unit"] = canonical_unit
    return result, station_rows


def write_csv(path: Path, rows: list[dict], fieldnames: list[str] | None = None):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    fieldnames = fieldnames or list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def markdown_summary(inventory: list[dict], quality: list[dict], station_rows: list[dict]) -> str:
    lines = [
        "# HIMPIL source-data inventory",
        "",
        "Generated by `python/inspect_data.py`. Source files are unchanged.",
        "",
        "## File and worksheet inventory",
        "",
        "| File | Sheet | Variable | Frequency | Period | Stations | Numeric cells | Missing % | Source unit |",
        "|---|---|---|---|---|---:|---:|---:|---|",
    ]
    for row in inventory:
        period = f"{row['year_start']}-{row['year_end']}" if row["year_start"] else ""
        lines.append(
            f"| {row['file']} | {row['sheet']} | {row['variable']} | {row['frequency']} | "
            f"{period} | {row['station_count']} | {row['numeric_cells']} | "
            f"{row['missing_percent']} | {row['source_unit']} |"
        )

    lines += [
        "",
        "## Data-quality flags",
        "",
        "- `...` is not defined in the source notes for most workbooks; it is treated as missing, not zero.",
        "- `**` means the station is no longer monitoring (water-quality workbooks).",
        "- The replacement/no-data symbol is treated as missing, not zero.",
        "- 2025 rainfall and wind are provisional and incomplete. Seven stations have all 12 months;",
        "  the other stations have January-June values only. The source says remaining synoptic data",
        "  are still undergoing verification.",
        "- Water-quality files contain annual station summaries, not individual sampling dates.",
        "- No workbook contains fish kill events or a fish kill label.",
        "",
        "## Lake stations found",
        "",
    ]
    lake = sorted({row["station"] for row in station_rows
                   if row["station_group"] == "Laguna de Bay (lake)"})
    for name in lake:
        lines.append(f"- {name}")

    lines += [
        "",
        "## Files that still need to be supplied",
        "",
        "- `HIMPIL.docx` was copied from Downloads into `data/raw/source_documents/`; it confirms "
        "the intended inputs but is a proposal, not a data source.",
        "- A dated historical fish kill incident dataset was not present.",
        "- Daily/dated water-quality observations were not present in the monitoring workbooks.",
        "  The only dated measurements found are the 2019 UPD article's during-event sample.",
        "",
    ]
    return "\n".join(lines)


def main():
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    workbooks = sorted(RAW_DIR.glob("*.xlsx"))
    if not workbooks:
        raise SystemExit(f"No .xlsx files found in {RAW_DIR}")

    inventory = []
    quality = []
    station_rows = []
    for path in workbooks:
        variable, unit, low, high = variable_info(path.name)
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
        for ws in workbook.worksheets:
            row, stations = inspect_sheet(ws, path.name, variable, unit, low, high)
            inventory.append(row)
            station_rows.extend(stations)
            quality.append({
                "file": row["file"],
                "sheet": row["sheet"],
                "variable": row["variable"],
                "total_station_period_cells": row["total_station_period_cells"],
                "numeric_cells": row["numeric_cells"],
                "missing_cells": row["missing_cells"],
                "missing_percent": row["missing_percent"],
                "ellipsis_undefined_cells": row["ellipsis_undefined_cells"],
                "no_data_symbol_cells": row["no_data_symbol_cells"],
                "station_discontinued_cells": row["station_discontinued_cells"],
                "blank_cells": row["blank_cells"],
                "other_non_numeric_cells": row["other_non_numeric_cells"],
                "out_of_range_cells": row["out_of_range_cells"],
            })
        workbook.close()

    write_csv(TABLE_DIR / "dataset_inventory.csv", inventory)
    write_csv(TABLE_DIR / "data_quality_report.csv", quality)
    write_csv(TABLE_DIR / "station_inventory.csv", station_rows)
    (TABLE_DIR / "dataset_inventory.md").write_text(
        markdown_summary(inventory, quality, station_rows), encoding="utf-8"
    )

    print("HIMPIL data inspection complete")
    print(f"Workbooks inspected : {len(workbooks)}")
    print(f"Worksheets inspected: {len(inventory)}")
    print(f"Station rows found  : {len(station_rows)}")
    print(f"Outputs             : {TABLE_DIR}")
    print("\nFile-level summary:")
    for path in workbooks:
        rows = [r for r in inventory if r["file"] == path.name]
        total_cells = sum(r["total_station_period_cells"] for r in rows)
        numeric = sum(r["numeric_cells"] for r in rows)
        missing = total_cells - numeric
        pct = round(100 * missing / total_cells, 2) if total_cells else 0
        variable = rows[0]["variable"] if rows else "Unknown"
        print(f"  {path.name}")
        print(f"    variable={variable}; sheets={len(rows)}; cells={total_cells}; "
              f"missing={missing} ({pct}%)")
    print("\nKey warnings:")
    print("  - Water-quality values are annual station summaries, not dated observations.")
    print("  - Rainfall/wind values are monthly station summaries; 2025 is incomplete and provisional.")
    print("  - No fish kill event file was found, so supervised labels cannot be created yet.")
    print("  - '...' is undefined in the source notes and is treated as missing.")


if __name__ == "__main__":
    main()
